# Junta a extração dos guias (extracao/g*.json) + QR (qr_coords.json), geocodifica o que faltar, classifica rotas/categorias e gera a planilha turismo_rural_rj.xlsx e atrativos.json
import json, math, re, time, unicodedata, urllib.request, urllib.parse, os
from collections import defaultdict, Counter
import sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
BASE = Path(__file__).resolve().parent   # dados-fonte/ — os caminhos não dependem de onde o script é chamado
RAIZ = BASE.parent
W = str(BASE) + '/'                      # qr_coords.json, geocache.json, atrativos.json (saída intermediária)
EXTR = W + 'extracao/'                   # g01..g12.json: atrativos lidos dos 12 guias (já com as correções da auditoria)
IBGE = W + 'ibge/'                       # malha municipal e nomes (IBGE)

def norm(s): return unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower().strip()

A = [x for f in range(1, 13) for x in json.load(open(f'{EXTR}g{f:02}.json', encoding='utf-8'))]
# nome oficial do IBGE para o município (o guia às vezes grafa diferente)
OFICIAL = {norm(m['nome']): m['nome'] for m in json.load(open(IBGE + 'nomes.json', encoding='utf-8'))}
ALIAS = {'armacao de buzios': 'armacao dos buzios', 'buzios': 'armacao dos buzios', 'paraty': 'paraty', 'parati': 'paraty',
         'trajano de morais': 'trajano de moraes', 'sao jose do vale do rio preto': 'sao jose do vale do rio preto'}
for a in A:
    n = norm(a['municipio']); n = ALIAS.get(n, n)
    if n in OFICIAL: a['municipio'] = OFICIAL[n]
qr = json.load(open(W + 'qr_coords.json'))

# 1) casar QR por página: mesma quantidade -> ordem de leitura (y, depois x)
qp = defaultdict(list)
for q in qr: qp[(q['pdf'], q['page'])].append(q)
ap = defaultdict(list)
for a in A:
    a['page'] = int(re.sub(r'\D', '', str(a['page'])) or 0)  # extratores gravaram ora 4, ora "p004"
    a.update(lat=None, lng=None, precisao='', mapa_url='')
    if a.get('tem_qr'): ap[(a['pdf'], a['page'])].append(a)
for k, items in ap.items():
    qs = sorted(qp.get(k, []), key=lambda q: (round(q['y'], 1), q['x']))
    items.sort(key=lambda a: (a['coluna'] != 'esq', a['ordem']))
    if len(qs) == len(items):  # ponytail: casamento só por contagem; páginas com QR não lido caem no geocoder
        for a, q in zip(items, qs):
            a['mapa_url'] = q['final'] if q['final'].startswith('http') else q['url']
            if q['lat'] and -23.5 < q['lat'] < -20.7 and -45 < q['lng'] < -40.9:
                a.update(lat=q['lat'], lng=q['lng'], precisao='QR (Google Maps)')

# 2) centroides dos municípios (IBGE)
nomes = {str(m['id']): m['nome'] for m in json.load(open(IBGE + 'nomes.json', encoding='utf-8'))}
geo = json.load(open(IBGE + 'mun.geojson', encoding='utf-8'))
cent = {}
for ft in geo['features']:
    pts = [p for poly in ft['geometry']['coordinates'] for p in poly[0]] if ft['geometry']['type'] == 'MultiPolygon' else ft['geometry']['coordinates'][0]
    n = nomes[ft['properties']['codarea']]; ft['properties']['nome'] = n
    cent[norm(n)] = (sum(p[1] for p in pts) / len(pts), sum(p[0] for p in pts) / len(pts))

# 2b) coordenada do QR só vale se (a) a própria URL do Google Maps traz a coordenada (!3d..!4d.., @lat,lng ou q=lat,lng) —
# links de "lugar" sem coordenada deixavam a extração cair no centro padrão do mapa (Rio/Niterói), igual para todos — e
# (b) o ponto cai a menos de 12 km do município declarado (casamento QR↔atrativo trocado ou lugar homônimo em outro canto)
def coord_da_url(u):
    u = urllib.parse.unquote(u)
    m = re.search(r'!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)', u) or re.search(r'@(-?\d+\.\d+),(-?\d+\.\d+)', u) or re.search(r'[?&](?:q|query|ll)=(-?\d+\.\d+),(-?\d+\.\d+)', u)
    return (float(m[1]), float(m[2])) if m else None
poligonos = {}
for ft in geo['features']:
    g = ft['geometry']; poligonos[norm(ft['properties']['nome'])] = g['coordinates'] if g['type'] == 'MultiPolygon' else [g['coordinates']]
def dentro(x, y, anel):
    c = False; j = len(anel) - 1
    for i in range(len(anel)):
        xi, yi = anel[i]; xj, yj = anel[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi: c = not c
        j = i
    return c
def km_fora(lat, lng, mun):  # 0 se dentro; senão, distância aproximada ao vértice mais próximo
    ps = poligonos.get(norm(mun))
    if not ps: return 0
    if any(dentro(lng, lat, p[0]) for p in ps): return 0
    return min(math.hypot((lng - px) * 103, (lat - py) * 110.6) for p in ps for anel in p for px, py in anel)
rejeitados = []
for a in A:
    if a['precisao'] != 'QR (Google Maps)': continue
    cu = coord_da_url(a['mapa_url'])
    pl = re.search(r'/maps/place/([^/@?]+)', urllib.parse.unquote(a['mapa_url']))
    so_cidade = bool(pl) and re.fullmatch(re.escape(norm(a['municipio'])) + r'(, rj)?(, \d{5}(-\d{3})?)?', norm(pl[1].replace('+', ' ')).strip()) is not None
    motivo = 'QR aponta só para a cidade' if so_cidade else 'sem coordenada na URL' if not cu else (f'{km_fora(a["lat"], a["lng"], a["municipio"]):.0f} km fora do município' if km_fora(a['lat'], a['lng'], a['municipio']) > 12 else '')
    if motivo:
        rejeitados.append((a['municipio'], a['nome'], motivo))
        a.update(lat=None, lng=None, precisao='', mapa_url='')
print('QR rejeitados (vão para endereço/município):', len(rejeitados)); [print('  ', r) for r in rejeitados]
# mesma coordenada em 3+ atrativos = pino da localidade (bairro), não do estabelecimento: vale como ponto aproximado, e o link busca pelo nome
repetida = Counter((round(a['lat'], 4), round(a['lng'], 4)) for a in A if a['precisao'] == 'QR (Google Maps)')
for a in A:
    if a['precisao'] == 'QR (Google Maps)' and repetida[(round(a['lat'], 4), round(a['lng'], 4))] >= 3:
        a.update(precisao='QR (localidade)', mapa_url='')

# 3) geocoder (Nominatim, 1 req/s, cache em disco)
cache_f = W + 'geocache.json'
cache = json.load(open(cache_f, encoding='utf-8')) if os.path.exists(cache_f) else {}
def nominatim(q):
    if q not in cache:
        url = 'https://nominatim.openstreetmap.org/search?' + urllib.parse.urlencode(dict(q=q, format='json', limit=1, countrycodes='br'))
        try: cache[q] = json.load(urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'prototipo-mapa-turismo-rural-rj'}), timeout=30))
        except Exception: cache[q] = []
        time.sleep(1.1)
        if len(cache) % 20 == 0: json.dump(cache, open(cache_f, 'w', encoding='utf-8'))
    r = cache[q]
    return (float(r[0]['lat']), float(r[0]['lon'])) if r else None

for a in A:
    if a['lat']: continue
    c = cent.get(norm(a['municipio']))
    hit = a['endereco'] and nominatim(f"{a['endereco']}, {a['municipio']}, RJ")
    # aceita só se cair perto do município (evita homônimos em outro canto)
    if hit and c and abs(hit[0] - c[0]) < 0.35 and abs(hit[1] - c[1]) < 0.35:
        a.update(lat=hit[0], lng=hit[1], precisao='endereço (aproximado)')
    elif c:
        a.update(lat=c[0], lng=c[1], precisao='centro do município')
json.dump(cache, open(cache_f, 'w', encoding='utf-8'))

# link do mapa: só vale se for do Google Maps (alguns QR levam a página expirada do me-qr, a uma busca ou a um site);
# senão, busca pela coordenada (quando é real) ou pelo nome
for a in A:
    if not re.search(r'google\.[a-z.]+/maps|maps\.app\.goo\.gl|goo\.gl/maps', a['mapa_url']):
        q = f"{a['lat']},{a['lng']}" if a['precisao'] == 'QR (Google Maps)' else f"{a['nome']}, {a['municipio']}, RJ"  # ponto aproximado: melhor buscar pelo nome
        a['mapa_url'] = 'https://www.google.com/maps/search/?api=1&query=' + urllib.parse.quote(q)

# 4) categorias (produtos) e experiências por palavra-chave
CAT = {
 'Café': r'\bcaf[eé]\b(?! da manh[aã]| colonial)|cafeicul|cafezal|torrefa', 'Cachaça e bebidas': r'cacha[cç]a|alambique|cervej|vinho|vin[ií]cola|licor|destil',
 'Queijos e laticínios': r'queij|latic[ií]n|iogurte|leite', 'Doces e compotas': r'\bdoce|compota|geleia|goiabada|rapadura',
 'Mel e abelhas': r'\bmel\b|apic|apiár|abelha|meliponi', 'Frutas': r'fruta|pomar|morango|banana|laranja|caqui|\buvas?\b|lim[aã]o|maracuj|jabuticab',
 'Flores e plantas': r'\bflor|orqu[ií]d|plantas|paisagis|viveiro|horto', 'Hortaliças e orgânicos': r'hortali|org[aâ]nic|horta\b|agroecol',
 'Gastronomia': r'restaurante|gastron|culin[aá]ri|comida|cozinha|almo[cç]o|caf[eé] colonial|fog[aã]o a lenha',
 'Hospedagem': r'pousada|hospedag|chal[eé]|hotel|camping|su[ií]te', 'Natureza e aventura': r'trilha|cachoeira|ecoturis|mata atl|observa[cç][aã]o de aves|rapel|tirolesa|po[cç]o',
 'Fazenda histórica': r'hist[oó]ric|(?<!caf[eé] )colonial|s[eé]culo|imperial|\bbar[aã]o\b|senzala', 'Artesanato': r'artesan|ateli|cer[aâ]mica|tecelag',
 'Pesca e criação': r'pesque|pesca|cria[cç][aã]o|cavalo|equita|haras|fazendinha|animais',
}
EXP = {
 'Degustação': r'degusta', 'Colheita': r'\bcolh(a|er|eita)', 'Trilha / natureza': r'trilha|cachoeira|banho de rio|po[cç]o',
 'Visita guiada': r'visita guiada|visita[cç][aã]o|tour|conhecer o processo|processo de produ', 'Compra direta': r'loja|venda|vend(e|em)|comprar|empório|feira',
 'Com crianças': r'crian[cç]a|fam[ií]lia|fazendinha|parquinho|pula-pula|infantil', 'Cavalgada': r'cavalg|equita|passeio a cavalo',
}
for a in A:
    t = (a['nome'] + ' ' + a['descricao']).lower()
    a['categorias'] = '; '.join(k for k, v in CAT.items() if re.search(v, t)) or 'Outros'
    a['experiencias'] = '; '.join(k for k, v in EXP.items() if re.search(v, t))
    a['categoria_principal'] = a['categorias'].split('; ')[0]

# rotas temáticas (prioridade: a primeira que casar vira a rota principal / cor no telão)
ROTAS = [
 ('Café e Fazendas Históricas', r'\bcaf[eé]\b(?! da manh[aã]| colonial)|cafeicul|cafezal|torrefa|fazenda hist|fazenda centen|s[eé]culo xix|\bbar[aã]o\b|\bimperial\b'),
 ('Queijos e Sabores da Montanha', r'queij|latic[ií]n|doce de leite|manteiga|iogurte'),
 ('Cachaças, Cervejas e Doces', r'cacha[cç]a|alambique|cervej|vin[ií]cola|\bvinhos?\b|licor|\bdoces?\b|compota|geleia|\bmel\b|apic|agroind|casa da farinha|\bfarinha\b'),
 ('Frutas, Flores e Colheitas', r'fruta|pomar|morango|\bflor(es)?\b|floricult|orqu[ií]d|\bcolh(a|er|eita)|hortali|\bhortas?\b|org[aâ]nic|agroecol|plantas ornamentais'),
 ('Águas e Pesca', r'pesque|\bpesca|pescador|cai[cç]ara|frutos do mar|mariscos|\bostras?\b'),
 ('Mata Atlântica e Natureza', r'mata atl[aâ]ntica|agroflorest|pupunha|palmito|trilha|cachoeira|observa[cç][aã]o de aves|quilomb|ind[ií]gena|ecoturis|caminhada'),
 ('Mesa e Pouso Rural', r'restaurante|gastron|culin[aá]ri|comida|cozinha|almo[cç]o|fog[aã]o a lenha|caf[eé] da manh|caf[eé] colonial|pousada|hospedag|hotel|chal[eé]s?\b|camping'),
 ('Cultura, Artesanato e Festas', r'artesan|ateli[eê]|cer[aâ]mica|tecelag|\bfesta|folia|cavalhada|museu|\bigreja|capela|cultur'),
]
for a in A:
    t = (a['nome'] + ' ' + a['descricao']).lower()
    rs = [n for n, v in ROTAS if re.search(v, t)]
    a['rotas'] = '; '.join(rs); a['rota_principal'] = rs[0] if rs else 'Vivências no Campo'

# 5) Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
COLS = ['regiao', 'municipio', 'nome', 'endereco', 'rota_principal', 'rotas', 'categoria_principal', 'categorias', 'experiencias', 'funcionamento', 'telefone',
        'instagram', 'site', 'wifi', 'agendar', 'pet', 'lat', 'lng', 'precisao', 'mapa_url', 'descricao', 'pdf', 'page']
HEAD = ['Região', 'Município', 'Nome', 'Endereço', 'Rota principal', 'Rotas', 'Categoria principal', 'Categorias', 'Experiências', 'Funcionamento', 'Telefone',
        'Instagram', 'Site', 'Wi-Fi', 'Precisa agendar', 'Pet friendly', 'Latitude', 'Longitude', 'Precisão da localização', 'Link do mapa', 'Descrição', 'Guia nº', 'Página']
wb = Workbook(); ws = wb.active; ws.title = 'Atrativos'; ws.append(HEAD)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E3D')
for a in sorted(A, key=lambda a: (a['pdf'], a['municipio'], a['page'], a['ordem'])): ws.append([a.get(k) for k in COLS])
for i, w in enumerate([18, 22, 34, 36, 28, 40, 20, 30, 28, 26, 24, 20, 26, 7, 9, 9, 11, 11, 20, 30, 80, 7, 7], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions

# resumo por região x rota
rotas_ord = [n for n, _ in ROTAS] + ['Vivências no Campo']
regs = list(dict.fromkeys(a['regiao'] for a in sorted(A, key=lambda a: a['pdf'])))
cnt = Counter((a['regiao'], a['rota_principal']) for a in A)
rs = wb.create_sheet('Resumo'); rs.append(['Região'] + rotas_ord + ['Total', 'Com ponto exato (QR)', 'Municípios'])
for c in rs[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E3D')
for r in regs:
    da = [a for a in A if a['regiao'] == r]
    rs.append([r] + [cnt[(r, ro)] for ro in rotas_ord] + [len(da), sum(a['precisao'] == 'QR (Google Maps)' for a in da), len({a['municipio'] for a in da})])
rs.append(['Total'] + [sum(cnt[(r, ro)] for r in regs) for ro in rotas_ord] + [len(A), sum(a['precisao'] == 'QR (Google Maps)' for a in A), len({a['municipio'] for a in A})])
for c in rs[rs.max_row]: c.font = Font(bold=True)
rs.column_dimensions['A'].width = 24
for i in range(2, len(rotas_ord) + 5): rs.column_dimensions[get_column_letter(i)].width = 16
rs.freeze_panes = 'B2'

lm = wb.create_sheet('Leia-me')
for linha in [
    ['Turismo Rural RJ — atrativos extraídos dos 12 guias (protótipo)'],
    ['Fonte', 'Guias "Turismo Rural RJ" (TurisRio / Setur-RJ), https://www.rj.gov.br/turisrio/node/132 — 12 PDFs, um por região turística.'],
    ['Extração', 'Texto lido das páginas (muitas são só imagem) por leitura assistida, auditado guia a guia contra as páginas; cada correção foi confirmada por um segundo verificador.'],
    ['Coordenadas', 'Precisão "QR (Google Maps)": coordenada do lugar, lida do QR "Ver mapa" impresso no guia (só aceita se a URL do QR traz a coordenada e o ponto cai no município). "QR (localidade)": o QR leva ao bairro/localidade, com a mesma coordenada para vários atrativos (aproximado). "endereço (aproximado)": endereço geocodificado no OpenStreetMap/Nominatim. "centro do município": o guia não traz coordenada utilizável; lat/long indicam só o centro do município.'],
    ['Rotas', 'Rota temática atribuída automaticamente por palavras-chave do nome e da descrição (a primeira que casar é a rota principal). São rotas sugeridas de promoção, não rotas oficiais.'],
    ['Categorias / Experiências', 'Também por palavras-chave; servem para filtro e não substituem a curadoria com os empreendimentos.'],
    ['Cuidados', 'Horários, telefones e agendamento podem ter mudado desde a publicação dos guias; confirme com o empreendimento antes de divulgar.'],
    ['Corte', '15/10/2026'],
]: lm.append(linha)
lm['A1'].font = Font(bold=True, size=13); lm.column_dimensions['A'].width = 26; lm.column_dimensions['B'].width = 120
wb.save(str(RAIZ / 'turismo_rural_rj.xlsx'))

from collections import Counter
print(len(A), 'atrativos |', Counter(a['precisao'] for a in A))
print(Counter(a['categoria_principal'] for a in A).most_common())
print(Counter(a['rota_principal'] for a in A).most_common())

json.dump(A, open(W + 'atrativos.json', 'w', encoding='utf-8'), ensure_ascii=False)
