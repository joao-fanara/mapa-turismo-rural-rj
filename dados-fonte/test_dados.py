"""Testes dos dados gerados (dados-fonte -> telao/dados e planilha). Rodar: python -m unittest dados-fonte/test_dados.py -v

Conferem invariantes que já falharam de verdade nesta base — cada uma é teste de regressão de um bug encontrado:
coordenada padrão do Google Maps repetida em vários municípios, QR que aponta só para a cidade, ponto fora do município.
"""
import json
import math
import re
import unicodedata
import unittest
from collections import Counter
from pathlib import Path

BASE = Path(__file__).resolve().parent
RAIZ = BASE.parent
DADOS = RAIZ / 'telao' / 'dados'
norm = lambda s: unicodedata.normalize('NFKD', s or '').encode('ascii', 'ignore').decode().lower().strip()
carrega = lambda p: json.loads(Path(p).read_text(encoding='utf-8'))

REG, FIC, LEN, MUN = (carrega(DADOS / f) for f in ('registros.json', 'fichas.json', 'lentes.json', 'municipios-3d.json'))
ATR = carrega(BASE / 'atrativos.json')
NOMES = {norm(m['nome']) for m in carrega(BASE / 'ibge' / 'nomes.json')}
GEO = carrega(BASE / 'ibge' / 'mun.geojson')
NOME_POR_COD = {str(m['id']): m['nome'] for m in carrega(BASE / 'ibge' / 'nomes.json')}
POLIGONOS = {}
for ft in GEO['features']:
    g = ft['geometry']
    POLIGONOS[norm(NOME_POR_COD[ft['properties']['codarea']])] = g['coordinates'] if g['type'] == 'MultiPolygon' else [g['coordinates']]
PL = REG['meta']['plano_local']
REGIOES_TURISTICAS = {'Vale do Café', 'Costa do Sol', 'Agulhas Negras', 'Serra Verde Imperial', 'Região Metropolitana', 'Costa Verde', 'Costa Doce',
                      'Caminhos Coloniais', 'Caminhos da Serra', 'Caminhos da Mata', 'Baixada Verde', 'Águas do Noroeste'}
COORD_PADRAO_GOOGLE = {(-22.9016, -43.1686), (-22.8917, -43.1161)}  # vista padrão do Google Maps (Rio/Niterói) quando o QR não acha o lugar


def dentro(x, y, anel):
    c, j = False, len(anel) - 1
    for i in range(len(anel)):
        (xi, yi), (xj, yj) = anel[i], anel[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            c = not c
        j = i
    return c


def km_fora(lat, lng, municipio):
    ps = POLIGONOS[norm(municipio)]
    if any(dentro(lng, lat, p[0]) for p in ps):
        return 0.0
    return min(math.hypot((lng - px) * 103, (lat - py) * 110.6) for p in ps for anel in p for px, py in anel)


class Consistencia(unittest.TestCase):
    def test_totais_batem_em_todos_os_arquivos(self):
        n = len(REG['registros'])
        self.assertEqual(n, REG['meta']['total'])
        self.assertEqual(n, len(FIC))
        self.assertEqual(n, len(ATR))
        self.assertEqual(n, len(LEN['por_registro']))
        self.assertEqual(n, sum(a['total'] for a in REG['agregados']['por_municipio'].values()))
        self.assertEqual(n, sum(g['total'] for g in REG['grupos']))

    def test_codigos_unicos_e_todo_registro_tem_ficha(self):
        cods = [r[5] for r in REG['registros']]
        self.assertEqual(len(cods), len(set(cods)))
        self.assertEqual(set(cods), set(FIC))

    def test_todo_atrativo_tem_nome_e_municipio_oficial_e_regiao_turistica(self):
        for a in ATR:
            self.assertTrue(a['nome'].strip(), a)
            self.assertIn(norm(a['municipio']), NOMES, a['municipio'])
            self.assertIn(a['regiao'], REGIOES_TURISTICAS, a['nome'])

    def test_rota_principal_existe_nos_grupos(self):
        slugs = {g['slug'] for g in REG['grupos']}
        self.assertTrue(all(f['rota_principal'] in slugs for f in FIC.values()))

    def test_municipios_do_telao_tem_regiao_turistica_ou_outros(self):
        for m in MUN['municipios']:
            self.assertTrue(m['regiao'] in REGIOES_TURISTICAS or m['regiao'] == 'Outros municípios', m['nome'])
        self.assertGreater(len(MUN['divisas_regioes']), 100)

    def test_texto_sem_lixo(self):
        for c, f in FIC.items():
            for campo in ('nome', 'endereco', 'descricao', 'funcionamento', 'telefone'):
                self.assertNotRegex(f[campo] or '', r'undefined|NaN|�', f'{c}.{campo}')


class Localizacao(unittest.TestCase):
    def test_classes_de_precisao_e_regra_do_ponto_municipal(self):
        permitido = {p['classe']: p['ponto_exato_permitido'] for p in REG['precisoes']}
        self.assertEqual(permitido, {'endereco': True, 'referencia': True, 'municipal': False})
        for r in REG['registros']:
            f = FIC[r[5]]
            if r[3] == 'municipal':
                self.assertIsNone(f['lat'], r[5])  # sem coordenada inventada: nada de centroide na ficha
            else:
                self.assertIsNotNone(f['lat'], r[5])

    def test_plano_km_e_o_inverso_da_coordenada(self):
        for r in REG['registros']:
            f = FIC[r[5]]
            if f['lat'] is None:
                continue
            lon = PL['lon0'] + r[0] / PL['km_por_grau_lon']
            lat = PL['lat0'] + r[1] / PL['km_por_grau_lat']
            self.assertAlmostEqual(lat, f['lat'], places=2, msg=r[5])
            self.assertAlmostEqual(lon, f['lng'], places=2, msg=r[5])

    def test_pontos_dentro_do_estado(self):
        for f in FIC.values():
            if f['lat'] is not None:
                self.assertTrue(-23.4 < f['lat'] < -20.7 and -45.1 < f['lng'] < -40.9, f['nome'])

    def test_regressao_coordenada_padrao_do_google_nunca_vira_ponto(self):
        for f in FIC.values():
            if f['lat'] is not None:
                self.assertNotIn((round(f['lat'], 4), round(f['lng'], 4)), COORD_PADRAO_GOOGLE, f['nome'])

    def test_regressao_ponto_exato_nao_se_repete_entre_atrativos(self):
        exatos = Counter((round(FIC[r[5]]['lat'], 4), round(FIC[r[5]]['lng'], 4)) for r in REG['registros'] if r[3] == 'endereco')
        self.assertEqual({k: v for k, v in exatos.items() if v >= 3}, {})  # 3+ atrativos no mesmo ponto = pino do bairro, não do lugar

    def test_regressao_ponto_exato_cai_perto_do_municipio_declarado(self):
        longe = [(f['nome'], f['municipio'], round(km_fora(f['lat'], f['lng'], f['municipio']))) for r in REG['registros'] if r[3] == 'endereco'
                 for f in [FIC[r[5]]] if km_fora(f['lat'], f['lng'], f['municipio']) > 12]
        self.assertEqual(longe, [])

    def test_link_do_mapa_e_sempre_google_maps(self):
        for f in FIC.values():
            self.assertRegex(f['mapa_url'], r'^https://www\.google\.(com|com\.br)/maps', f['nome'])

    def test_regressao_link_nao_e_so_a_cidade(self):
        for f in FIC.values():
            m = re.search(r'/maps/place/([^/@?]+)', f['mapa_url'])
            if m:
                lugar = norm(re.sub(r'\+', ' ', m[1]))
                self.assertIsNone(re.fullmatch(re.escape(norm(f['municipio'])) + r'(, rj)?(, \d{5}(-\d{3})?)?', lugar), f['nome'])


class Planilha(unittest.TestCase):
    def test_planilha_tem_as_abas_e_todos_os_atrativos(self):
        from openpyxl import load_workbook  # pip install openpyxl
        wb = load_workbook(RAIZ / 'turismo_rural_rj.xlsx', read_only=True)
        self.assertEqual(wb.sheetnames, ['Atrativos', 'Resumo', 'Leia-me'])
        self.assertEqual(wb['Atrativos'].max_row, len(ATR) + 1)


if __name__ == '__main__':
    unittest.main()
